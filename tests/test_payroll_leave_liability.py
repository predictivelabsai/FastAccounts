from datetime import date
from decimal import Decimal
import pytest

from payroll_leave_liability import (
    LeaveLot, accrued_days, consume_fifo, expiry_date, liability_value,
)
from payroll_leave_service import PayrollLeaveLiabilityService, _lots


def test_entitlement_accrues_pro_rata_by_calendar_day():
    assert accrued_days(28, date(2026, 7, 1), date(2026, 12, 31), 2026) == Decimal('14.1151')


def test_expiry_policies_and_fifo_consumption():
    assert expiry_date(2024, '1_january_year_plus_2') == date(2026, 1, 1)
    assert expiry_date(2024, '31_december_year_plus_1') == date(2025, 12, 31)
    lots = [LeaveLot(Decimal('10'), date(2025, 1, 1), date(2027, 1, 1)),
            LeaveLot(Decimal('7'), date(2026, 1, 1), None)]
    remaining, consumed, shortage = consume_fifo(lots, Decimal('12'), date(2026, 10, 1))
    assert [days for _, days in consumed] == [Decimal('10'), Decimal('2')]
    assert remaining[0].days == Decimal('5')
    assert shortage == 0


def test_expired_lot_is_not_available():
    remaining, consumed, shortage = consume_fifo(
        [LeaveLot(Decimal('4'), date(2024, 1, 1), date(2025, 12, 31))], 2,
        date(2026, 1, 1))
    assert remaining == [] and consumed == [] and shortage == 2


def test_liability_adds_employer_social_costs():
    value = liability_value(Decimal('5'), Decimal('100'), date(2026, 10, 31))
    assert value['wage_liability'] == '500.00'
    assert value['social_tax'] == '165.00'
    assert value['employer_unemployment'] == '4.00'
    assert value['total_liability'] == '669.00'


def test_persisted_source_rebuild_consumes_oldest_lot_and_records_expiry():
    lots, movements = _lots(
        [{'effective_from': '2024-01-01', 'days_per_year': '28'}], [],
        [{'id': 'leave', 'absence_type': 'annual_leave', 'start_date': '2024-06-01', 'end_date': '2024-06-03'}],
        {'basic_expiry_policy': '31_december_year_plus_1'}, '2024-01-01', date(2026, 1, 1))
    assert sum((lot.days for lot in lots), Decimal(0)) == Decimal('28.0767')
    assert any(row['type'] == 'consumption' and row['days'] == '3' for row in movements)
    assert any(row['type'] == 'expired' and row['date'] == '2025-12-31' for row in movements)


def test_leave_settings_are_tenant_scoped_and_default_off(db, ee_org, uk_org):
    service = PayrollLeaveLiabilityService(db)
    assert service.settings(ee_org['id'])['reserve_posting_enabled'] == 0
    saved = service.save_settings(ee_org['id'], basic_expiry_policy='1_january_year_plus_2',
        extended_expiry_policy='31_december_year_plus_1', reserve_posting_enabled=True,
        actor='owner@example.test')
    assert saved['basic_expiry_policy'] == '1_january_year_plus_2'
    assert service.settings(uk_org['id'])['reserve_posting_enabled'] == 0


def test_report_rebuilds_tenant_scoped_lots_and_values_social_costs(db, ee_org, uk_org):
    from io import BytesIO
    from openpyxl import load_workbook
    from payroll import PayrollService
    from core_utils import new_id
    employee = PayrollService(db).save_employee(ee_org['id'], actor='owner@example.test',
        name='Leave Report Employee', gross_salary='2200', employment_start_date='2026-01-01')
    with db.transaction() as tx:
        tx.execute('INSERT INTO payroll_leave_entitlements(id,organisation_id,employee_id,effective_from,days_per_year,created_at) VALUES (?,?,?,?,?,?)',
            (new_id(), ee_org['id'], employee['id'], '2026-01-01', '28', '2026-01-01'))
    report = PayrollLeaveLiabilityService(db).report(ee_org['id'], '2026-06-30')
    assert report['employees'][0]['remaining_days'] == '13.8849'
    assert Decimal(report['employees'][0]['social_tax']) > 0
    workbook = load_workbook(BytesIO(PayrollLeaveLiabilityService(db).workbook_xlsx(ee_org['id'], '2026-06-30')), data_only=True)
    assert workbook.sheetnames == ['Leave liability', 'Movements']
    assert workbook['Leave liability']['A2'].value == 'Leave Report Employee'
    assert workbook['Movements']['B2'].value == 'accrual'
    assert PayrollLeaveLiabilityService(db).report(uk_org['id'], '2026-06-30')['employees'] == []


def test_monthly_reserve_posts_balanced_once_and_reverses(db, ee_org):
    from payroll import PayrollService, PayrollConflict
    from core_utils import new_id
    service = PayrollLeaveLiabilityService(db)
    employee = PayrollService(db).save_employee(ee_org['id'], actor='owner@example.test',
        name='Reserve Employee', gross_salary='2200', employment_start_date='2026-01-01')
    with db.transaction() as tx:
        tx.execute('INSERT INTO payroll_leave_entitlements(id,organisation_id,employee_id,effective_from,days_per_year,created_at) VALUES (?,?,?,?,?,?)',
            (new_id(), ee_org['id'], employee['id'], '2026-01-01', '28', '2026-01-01'))
    service.save_settings(ee_org['id'], basic_expiry_policy='none', extended_expiry_policy='none',
        reserve_posting_enabled=True, actor='owner@example.test')
    posted = service.post_reserve(ee_org['id'], '2026-06', actor='owner@example.test')
    with db.transaction() as tx:
        entries = tx.rows('SELECT debit,credit FROM gl_entries WHERE posting_batch_id=?', (posted['posting_batch_id'],))
    assert sum((Decimal(row['debit']) for row in entries), Decimal(0)) == sum((Decimal(row['credit']) for row in entries), Decimal(0))
    with pytest.raises(PayrollConflict):
        service.post_reserve(ee_org['id'], '2026-06', actor='owner@example.test')
    reversal = service.reverse_reserve(ee_org['id'], '2026-06', reversal_date='2026-06-30', actor='owner@example.test', reason='Correction')
    assert reversal['reverses_batch_id'] == posted['posting_batch_id']


def test_leave_settings_api_requires_csrf_and_owner_or_admin(db, ee_org, monkeypatch):
    from fastapi.testclient import TestClient
    from api_app import api, current_user, User
    from organisations import OrganisationService
    from web_app import app
    monkeypatch.setenv('FASTACCOUNTS_DB', db.path)
    monkeypatch.setenv('DB_URL', '')
    actor = {'email': 'owner@example.test'}
    api.dependency_overrides[current_user] = lambda: User(email=actor['email'])
    client = TestClient(app)
    url = f"/api/organisations/{ee_org['id']}/payroll/leave-settings"
    body = {'basic_expiry_policy':'none','extended_expiry_policy':'none','reserve_posting_enabled':False}
    try:
        assert client.put(url, json=body).status_code == 403
        token = client.get('/api/csrf').json()['csrf_token']
        headers = {'X-CSRF-Token': token}
        assert client.put(url, json=body, headers=headers).status_code == 200
        OrganisationService(db).add_member(ee_org['id'], 'owner@example.test', 'accountant@example.test', 'accountant')
        actor['email'] = 'accountant@example.test'
        assert client.put(url, json=body, headers=headers).status_code == 403
    finally:
        api.dependency_overrides.clear()
