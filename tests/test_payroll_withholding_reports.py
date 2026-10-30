from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from payroll import PayrollService
from payroll_payouts import PayoutService
from payroll_withholdings import PayrollWithholdingService


ACTOR = 'owner@example.test'


def test_withholding_reports_group_and_total_by_employee_recipient_type(db, ee_org, uk_org):
    payroll = PayrollService(db)
    employee = payroll.save_employee(ee_org['id'], actor=ACTOR, name='Report employee', gross_salary='2000')
    service = PayrollWithholdingService(db)
    recipient = service.save_recipient(ee_org['id'], actor=ACTOR, name='Report recipient',
        registry_code='report-001', iban='EE382200221020145685')
    service.save_agreement(ee_org['id'], actor=ACTOR, employee_id=employee['id'], recipient_id=recipient['id'],
        component_type='union_fee', method='fixed_sum', start_date='2026-01-01', monthly_amount='12.50')
    run = payroll.create_run(ee_org['id'], period='2026-10', actor=ACTOR)
    payroll.approve(ee_org['id'], run['id'], actor=ACTOR)
    report = service.report(ee_org['id'], start='2026-10', end='2026-10')
    assert report['total'] == '12.50'
    assert report['rows'] == [{'employee_name': 'Report employee', 'recipient_name': 'Report recipient',
                               'component_type': 'union_fee', 'amount': '12.50'}]
    assert service.report(uk_org['id'], start='2026-10', end='2026-10')['rows'] == []
    workbook = load_workbook(BytesIO(service.report_xlsx(ee_org['id'], start='2026-10', end='2026-10')),
                             data_only=True)
    sheet = workbook['Withholdings']
    assert sheet['D4'].value == 12.5


def test_withholding_payout_report_totals_and_xlsx(db, ee_org):
    payroll = PayrollService(db)
    employee = payroll.save_employee(ee_org['id'], actor=ACTOR, name='Paid employee', gross_salary='2000')
    service = PayrollWithholdingService(db)
    recipient = service.save_recipient(ee_org['id'], actor=ACTOR, name='Payout recipient',
        registry_code='payout-001', iban='EE382200221020145685')
    service.save_agreement(ee_org['id'], actor=ACTOR, employee_id=employee['id'], recipient_id=recipient['id'],
        component_type='agreed_deduction', method='fixed_sum', start_date='2026-01-01', monthly_amount='18')
    run = payroll.create_run(ee_org['id'], period='2026-10', actor=ACTOR)
    payroll.approve(ee_org['id'], run['id'], actor=ACTOR)
    PayoutService(db).create(ee_org['id'], run['id'], actor=ACTOR, payout_type='withholding')
    report = service.payout_report(ee_org['id'], start='2026-10', end='2026-10')
    assert Decimal(report['total']) == Decimal('18.00')
    assert report['rows'][0]['status'] == 'Draft'
    workbook = load_workbook(BytesIO(service.report_xlsx(ee_org['id'], start='2026-10', end='2026-10', payouts=True)),
                             data_only=True)
    sheet = workbook['Withholding payouts']
    assert sheet['F4'].value == 18
