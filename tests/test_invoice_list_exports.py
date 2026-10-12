"""Filtered document lists, safe downloads and offline workspace controls."""
from __future__ import annotations

import csv
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from database import reset_database_cache
from documents import DocumentService
from organisations import OrganisationService

ACTOR = "owner@example.test"


@pytest.fixture(params=[("uk_org", "invoices"), ("uk_org", "bills"),
                        ("ee_org", "invoices"), ("ee_org", "bills")])
def listing(request, db, uk_org, ee_org, monkeypatch):
    org = request.getfixturevalue(request.param[0])
    other = ee_org if org["id"] == uk_org["id"] else uk_org
    kind = request.param[1]
    service = DocumentService(db)

    def contact(book, name):
        return service.create_contact(book["id"], name=name, country_code=book["country_code"])

    primary = contact(org, "Willow, Design OÜ")
    secondary = contact(org, "Other synthetic contact")
    foreign = contact(other, "Foreign synthetic contact")

    def create(book, party, day, reference, currency=None, posted=False):
        account = db.scalar("SELECT id FROM accounts WHERE organisation_id=? AND system_role=?",
                            (book["id"], "SALES" if kind == "invoices" else "EXPENSE"))
        fields = dict(contact_id=party["id"], reference=reference, currency=currency,
                      actor=ACTOR, lines=[dict(account_id=account, unit_price="12.34")])
        if kind == "invoices":
            record = service.create_invoice(book["id"], issue_date=day, **fields)
            return service.issue_invoice(record["id"], actor=ACTOR) if posted else record
        record = service.create_bill(book["id"], supplier_number="SUP-" + reference,
                                     bill_date=day, due_date=day, **fields)
        if posted:
            service.submit_bill_for_review(record["id"], actor=ACTOR)
            return service.approve_bill(record["id"], actor=ACTOR)
        return record

    first = create(org, primary, "2026-01-10", "Needle 100%_", posted=True)
    second = create(org, secondary, "2026-02-10", "Second", currency="USD")
    third = create(org, primary, "2026-03-10", "Third")
    foreign_record = create(other, foreign, "2026-01-10", "Needle foreign", posted=True)
    OrganisationService(db).add_member(org["id"], ACTOR, "viewer@example.test", "viewer")
    monkeypatch.setenv("FASTACCOUNTS_DB", db.path)
    monkeypatch.setenv("DB_URL", "")
    monkeypatch.setenv("FASTACCOUNTS_ALLOW_TEST_AUTH", "true")
    reset_database_cache()
    from api_app import api
    with TestClient(api) as client:
        yield dict(client=client, root=f"/organisations/{org['id']}/{kind}", kind=kind,
                   org=org, first=first, second=second, third=third, primary=primary,
                   foreign=foreign, foreign_record=foreign_record, db=db)
    reset_database_cache()


def get(data, suffix="", params=None, email=ACTOR):
    response = data["client"].get(data["root"] + suffix, params=params, headers={"X-Test-User": email})
    assert response.status_code == 200, response.text
    return response


def test_each_filter_combined_search_and_tenant_isolation(listing):
    d = listing
    first, second, third = (d[key] for key in ("first", "second", "third"))
    ids = lambda params: {row["id"] for row in get(d, params=params).json()}
    assert ids({}) == {first["id"], second["id"], third["id"]}
    cases = [
        ({"status": first["status"]}, {first["id"]}),
        ({"status": "Draft"}, {second["id"], third["id"]}),
        ({"contact_id": d["primary"]["id"]}, {first["id"], third["id"]}),
        ({"date_from": "2026-02-10"}, {second["id"], third["id"]}),
        ({"date_to": "2026-02-10"}, {first["id"], second["id"]}),
        ({"currency": "USD"}, {second["id"]}),
        ({"q": "nEeDlE"}, {first["id"]}),
        ({"q": "willow"}, {first["id"], third["id"]}),
        ({"q": "OÜ"}, {first["id"], third["id"]}),
        ({"q": "100%_"}, {first["id"]}),
        ({"q": first.get("number") or first["supplier_number"]}, {first["id"]}),
        ({"status": first["status"], "contact_id": d["primary"]["id"],
          "date_from": "2026-01-10", "date_to": "2026-01-10",
          "currency": d["org"]["base_currency"], "q": "needle"}, {first["id"]}),
        ({"contact_id": d["foreign"]["id"]}, set()),
        ({"q": "foreign"}, set()),
        ({"status": "nonsense"}, set()),
        ({"currency": "nonsense"}, set()),
        ({"contact_id": "missing"}, set()),
        ({"q": "' OR 1=1 --"}, set()),
    ]
    for params, expected in cases:
        assert ids(params) == expected, params
    if d["kind"] == "invoices":
        assert ids({"status": "Overdue"}) == {first["id"]}


def test_csv_and_excel_use_same_filters_and_preserve_money(listing):
    d = listing
    number_key = "number" if d["kind"] == "invoices" else "supplier_number"
    for params in ({}, {"status": d["first"]["status"]},
                   {"contact_id": d["primary"]["id"], "date_to": "2026-01-10", "q": "needle"},
                   {"currency": "USD"}, {"contact_id": d["foreign"]["id"]},
                   {"status": "missing"}):
        selected = get(d, params=params).json()
        response = get(d, "/export.csv", params)
        assert response.headers["content-type"].startswith("text/csv")
        assert response.headers["content-disposition"] == f'attachment; filename="{d["kind"]}.csv"'
        csv_rows = list(csv.DictReader(StringIO(response.content.decode("utf-8-sig"))))
        assert len(csv_rows) == len(selected)
        assert [r["Number"] for r in csv_rows] == [r.get(number_key) or "" for r in selected]
        response = get(d, "/export.xlsx", params)
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
        book = load_workbook(BytesIO(response.content), data_only=False)
        sheet = book.active
        assert book.sheetnames == ["Invoices" if d["kind"] == "invoices" else "Bills"]
        assert sheet.max_row == len(selected) + 1
        if d["kind"] == "invoices":
            legacy = d["client"].get(d["root"].rsplit("/", 1)[0] + "/reports/invoices.xlsx",
                                      params=params, headers={"X-Test-User": "viewer@example.test"})
            assert legacy.status_code == 200
            legacy_sheet = load_workbook(BytesIO(legacy.content)).active
            # The established report workbook retains its final total row.
            assert legacy_sheet.max_row == len(selected) + 2
        for index, (source, exported) in enumerate(zip(selected, csv_rows), 2):
            assert exported["Contact name"] == source["contact_name"]
            assert exported["Reference"] == source["reference"]
            assert exported["Currency"] == source["currency"]
            assert exported["Gross"] == "12.34"
            assert sheet.cell(index, 3).value == source["contact_name"]
            assert Decimal(str(sheet.cell(index, 9).value)) == Decimal("12.34")
            assert sheet.cell(index, 9).number_format == "#,##0.00"
            assert sheet.cell(index, 9).data_type == "n"
            assert sheet.cell(index, 11).value == source["status"]
    # Read-only exports do not alter documents or their ledger batches.
    assert d["db"].one(f"SELECT * FROM {d['kind']} WHERE id=?", (d["first"]["id"],))["posting_batch_id"] == d["first"]["posting_batch_id"]


def test_invalid_dates_fail_cleanly_on_lists_and_exports(listing):
    d = listing
    for suffix in ("", "/export.csv", "/export.xlsx"):
        for params in ({"date_from": "yesterday"}, {"date_to": "2026-02-30"},
                       {"date_from": "20260101"}, {"date_to": "2026-1-01"},
                       {"date_from": "2026-03-01", "date_to": "2026-02-01"}):
            response = d["client"].get(d["root"] + suffix, params=params, headers={"X-Test-User": ACTOR})
            assert response.status_code == 400, response.text
        assert get(d, suffix, {"date_from": "2026-01-01", "date_to": "2026-01-01"}).content
    assert d["client"].get(d["root"] + "/export.pdf", headers={"X-Test-User": ACTOR}).status_code == 404


def test_viewer_downloads_and_nonmember_role_gates(listing):
    for suffix in ("", "/export.csv", "/export.xlsx"):
        get(listing, suffix, {"status": listing["first"]["status"]}, email="viewer@example.test")
        response = listing["client"].get(listing["root"] + suffix, headers={"X-Test-User": "intruder@example.test"})
        assert response.status_code == 403
        assert listing["client"].get(listing["root"] + suffix).status_code == 401


def test_tenant_join_and_formula_text_are_safe(listing):
    d = listing
    # The historical core allows a foreign contact ID; reads must still scope the join.
    with d["db"].transaction() as tx:
        tx.execute(f"UPDATE {d['kind']} SET contact_id=? WHERE id=?", (d["foreign"]["id"], d["second"]["id"]))
        tx.execute(f"UPDATE {d['kind']} SET reference=? WHERE id=?", (" =SUM(1,2)", d["third"]["id"]))
    assert d["second"]["id"] not in {r["id"] for r in get(d).json()}
    rows = list(csv.DictReader(StringIO(get(d, "/export.csv", {"q": "SUM"}).content.decode("utf-8-sig"))))
    assert rows[0]["Reference"] == "' =SUM(1,2)"
    sheet = load_workbook(BytesIO(get(d, "/export.xlsx", {"q": "SUM"}).content)).active
    assert sheet.cell(2, 2).value == " =SUM(1,2)"
    assert sheet.cell(2, 2).data_type == "s"


@pytest.mark.parametrize("language", ["en", "et", "lt", "lv"])
def test_workspace_filters_and_export_links(monkeypatch, language):
    playwright = pytest.importorskip("playwright.sync_api")
    from web import i18n
    from web_app import app
    monkeypatch.setenv("FASTACCOUNTS_DEFAULT_LANG", language)
    monkeypatch.setenv("FASTACCOUNTS_ALLOW_TEST_AUTH", "true")
    html = TestClient(app).get("/auth/test").text
    words = i18n.catalog(language)["workspace"]
    calls = []
    with playwright.sync_playwright() as runtime:
        browser = runtime.chromium.launch()
        page = browser.new_page()
        errors = []
        page.on("pageerror", lambda error: errors.append(str(error)))
        def respond(route):
            parsed = urlsplit(route.request.url)
            path = parsed.path
            if path == "/app":
                return route.fulfill(body=html, content_type="text/html")
            if path in ("/static/app.js", "/static/app.css"):
                return route.fulfill(path=str(Path(__file__).resolve().parents[1] / path.lstrip("/")))
            if path == "/api/organisations":
                return route.fulfill(json=[dict(id="org", name="Synthetic", country_code="EE", base_currency="EUR", role="viewer")])
            if path == "/api/csrf":
                return route.fulfill(json={"csrf_token": "synthetic"})
            if path.endswith("/contacts"):
                return route.fulfill(json=[dict(id="contact", name="Synthetic <contact>", contact_type="both")])
            if path.endswith(("/invoices", "/bills")):
                calls.append((path.rsplit("/", 1)[-1], parse_qs(parsed.query)))
                return route.fulfill(json=[], content_type="application/json")
            if path.startswith("/api/"):
                return route.fulfill(json=[], content_type="application/json")
            return route.fulfill(status=404)
        page.route("**/*", respond)
        page.goto("http://workspace.test/app#invoices")
        for kind in ("invoices", "bills"):
            if kind == "bills":
                page.locator("[data-view=bills]").click()
            control = lambda name: page.locator(f'[data-document-filter="{name}"]')
            playwright.expect(control("status")).to_have_value("")
            control("status").select_option("Draft")
            playwright.expect(control("status")).to_have_value("Draft")
            control("contact_id").select_option("contact")
            playwright.expect(control("contact_id")).to_have_value("contact")
            for name, value in (("date_from", "2026-01-01"), ("date_to", "2026-01-31"),
                                ("q", "Synthetic & <search>"), ("currency", "EUR")):
                control(name).fill(value)
                control(name).press("Tab")
                playwright.expect(control(name)).to_have_value(value)
                page.wait_for_function("document.querySelector('#app-content').getAttribute('aria-busy') === null")
            expected = dict(status=["Draft"], contact_id=["contact"], date_from=["2026-01-01"],
                            date_to=["2026-01-31"], q=["Synthetic & <search>"], currency=["EUR"])
            page.wait_for_function("document.querySelector('#app-content').getAttribute('aria-busy') === null")
            assert calls[-1] == (kind, expected)
            for extension, label in (("csv", "Export CSV"), ("xlsx", "Export Excel")):
                anchor = page.get_by_role("link", name=words[label], exact=True)
                href = urlsplit(anchor.get_attribute("href"))
                assert href.path.endswith(f"/{kind}/export.{extension}")
                assert parse_qs(href.query) == expected
            assert page.locator("#app-content contact, #app-content search").count() == 0
            page.set_viewport_size({"width": 390, "height": 844})
            assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
            page.set_viewport_size({"width": 1280, "height": 900})
        assert not errors
        browser.close()
