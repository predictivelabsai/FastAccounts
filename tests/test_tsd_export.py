from __future__ import annotations

import csv
import io
from decimal import Decimal
from xml.etree import ElementTree as ET

import pytest
from fastapi.testclient import TestClient

from connectors.providers import EMTATSDExportProvider
from payroll import PayrollService
from tsd import (
    TSD_AMOUNT_FIELDS,
    TSD_NS,
    TSD_PAYMENT_FIELDS,
    TSDExportService,
    XBRL_GL_CORE_NS,
)
from payroll_payday import PaydaySettingsService
from payroll_payouts import PayoutService


ACTOR = "owner@example.test"


def _approved_run(db, organisation):
    payroll = PayrollService(db)
    first = payroll.save_employee(
        organisation["id"], actor=ACTOR, name="Ada Arvutaja",
        personal_id="49001010018", gross_salary="2000",
        apply_tax_free_minimum=True,
    )
    payroll.save_employee(
        organisation["id"], actor=ACTOR, name="Bert Juhatuse liige",
        personal_id="38002020029", gross_salary="500", board_member=True,
    )
    run = payroll.create_run(organisation["id"], period="2026-10", actor=ACTOR)
    payroll.save_employee(
        organisation["id"], employee_id=first["id"], actor=ACTOR,
        personal_id="59912319999",
    )
    return payroll.approve(organisation["id"], run["id"], actor=ACTOR)


def _reviewed_preparation(db, organisation):
    run = _approved_run(db, organisation)
    service = TSDExportService(db)
    preparation = service.prepare(organisation["id"], run["id"], actor=ACTOR)
    assert preparation["status"] == "Review Required"
    with pytest.raises(ValueError, match="review"):
        service.export_xbrl_gl(organisation["id"], preparation["id"], actor=ACTOR)
    service.review(organisation["id"], preparation["id"], actor=ACTOR)
    return service, run, preparation


def test_xbrl_gl_has_one_complete_transaction_per_payment_and_reconciles(db, ee_org):
    service, run, preparation = _reviewed_preparation(db, ee_org)
    payload = service.export_xbrl_gl(ee_org["id"], preparation["id"], actor=ACTOR)
    root = ET.fromstring(payload)
    ns = {"gl": XBRL_GL_CORE_NS, "tsd": TSD_NS}
    headers = root.findall(".//gl:entryHeader", ns)
    assert len(headers) == len(run["items"]) == 2
    details = [header.find("gl:entryDetail", ns) for header in headers]
    assert all(detail is not None for detail in details)
    for detail in details:
        assert all(detail.find(f"tsd:{field}", ns) is not None for field in TSD_PAYMENT_FIELDS)
        assert detail.findtext("tsd:currency", namespaces=ns) == "EUR"
        assert detail.findtext("tsd:payment_date", namespaces=ns) == "2026-10-31"
        assert detail.findtext("tsd:filing_deadline", namespaces=ns) == "2026-11-10"
        assert detail.findtext("tsd:annex", namespaces=ns) == "1"
    personal_ids = {detail.findtext("tsd:recipient_personal_id", namespaces=ns) for detail in details}
    assert personal_ids == {"49001010018", "38002020029"}
    assert {detail.findtext("tsd:tsd_payment_type", namespaces=ns) for detail in details} == {"10", "21"}
    totals = {
        field: sum(
            (Decimal(detail.findtext(f"tsd:{field}", namespaces=ns)) for detail in details),
            Decimal("0"),
        )
        for field in TSD_AMOUNT_FIELDS
    }
    assert totals["gross_amount"] == Decimal(run["total_gross"])
    assert totals["net_amount"] == Decimal(run["total_net"])
    assert totals["employer_cost"] == Decimal(run["total_employer_cost"])
    assert root.findtext(".//tsd:annexScope", namespaces=ns) == "1-2"
    assert root.findtext(".//tsd:validationStatus", namespaces=ns) == "NOT_LIVE_VERIFIED_AGAINST_EMTA"


def test_transitional_csv_has_declared_fields_and_reconciles(db, ee_org):
    service, run, preparation = _reviewed_preparation(db, ee_org)
    payload = service.export_csv(ee_org["id"], preparation["id"], actor=ACTOR)
    rows = list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
    assert tuple(rows[0]) == TSD_PAYMENT_FIELDS
    assert len(rows) == len(run["items"])
    assert {row["payment_classification"] for row in rows} == {
        "EMPLOYMENT_SALARY", "BOARD_MEMBER_FEE"
    }
    assert {row["tsd_payment_type"] for row in rows} == {"10", "21"}
    assert sum((Decimal(row["gross_amount"]) for row in rows), Decimal("0")) == Decimal(run["total_gross"])
    assert sum((Decimal(row["net_amount"]) for row in rows), Decimal("0")) == Decimal(run["total_net"])
    assert sum((Decimal(row["employer_cost"]) for row in rows), Decimal("0")) == Decimal(run["total_employer_cost"])
    sync = db.one(
        "SELECT * FROM sync_runs WHERE organisation_id=? AND id=?",
        (ee_org["id"], preparation["id"]),
    )
    assert sync["direction"] == "export" and sync["object_type"] == "tsd"
    assert sync["status"] == "Completed" and sync["write_count"] == 2


def test_legacy_annex1_csv_golden_fields_and_format(db, ee_org):
    service, _, preparation = _reviewed_preparation(db, ee_org)
    payload = service.export_legacy_annex1_csv(ee_org['id'], preparation['id'], actor=ACTOR)
    assert payload.decode('utf-8') == (
        '1000;1010;1020;1030;1040;1050;1060;1070;1080;1090;1100;1110;1120;1130;1140;1150;1160;1170\r\n'
        '49001010018;Ada Arvutaja;10;2000.00;;;2000.00;0.00;0.00;0.00;660.00;40.00;2000.00;32.00;16.00;610;700.00;270.16\r\n'
        '38002020029;Bert Juhatuse liige;21;500.00;;;500.00;0.00;0.00;0.00;165.00;0.00;0.00;0.00;0.00;;0.00;110.00\r\n'
    )
    assert not payload.startswith(b'\xef\xbb\xbf')
    rows = list(csv.DictReader(io.StringIO(payload.decode('utf-8')), delimiter=';'))
    assert tuple(rows[0]) == ('1000','1010','1020','1030','1040','1050','1060','1070','1080','1090','1100','1110','1120','1130','1140','1150','1160','1170')
    assert len(rows) == 2
    assert rows[0]['1000'] == '49001010018' and rows[0]['1010'] == 'Ada Arvutaja'
    assert rows[0]['1020'] == '10' and rows[0]['1030'].endswith('.00')
    assert rows[0]['1040'] == rows[0]['1050'] == ''
    assert rows[1]['1020'] == '21' and rows[1]['1120'] == '0.00'
    assert all(len(row['1030'].rsplit('.', 1)[1]) == 2 for row in rows)


def test_legacy_annex1_csv_mixed_payment_type_golden(db, ee_org):
    payroll = PayrollService(db)
    payroll.save_employee(ee_org['id'], actor=ACTOR, name='Mixed CSV recipient',
                          personal_id='49001010018', gross_salary='2000', apply_tax_free_minimum=True)
    types = {row['preset_key']: row for row in payroll.component_types(ee_org['id'])}
    flags = dict(income_tax=True, social_tax=True, ui_employee=True, ui_employer=True,
                 funded_pension=True, social_minimum=False)
    board = payroll.save_component_type(ee_org['id'], actor=ACTOR, name='Board fee',
        direction='income', tsd_payment_type=21, **(flags | {'ui_employee': False,
        'ui_employer': False, 'social_minimum': False}))
    vacation = payroll.save_component_type(ee_org['id'], actor=ACTOR, name='Vacation pay',
        direction='income', tsd_payment_type=10, **flags)
    run = payroll.create_run(ee_org['id'], period='2026-10', actor=ACTOR)
    payroll.update_item(ee_org['id'], run['id'], run['items'][0]['id'], actor=ACTOR, components=[
        {'type_id': types['bonus']['id'], 'amount': '100'},
        {'type_id': vacation['id'], 'amount': '100'},
        {'type_id': board['id'], 'amount': '300'},
        {'type_id': types['contract']['id'], 'amount': '400'},
        {'type_id': types['employer_sick_pay']['id'], 'amount': '100'},
    ])
    approved = payroll.approve(ee_org['id'], run['id'], actor=ACTOR)
    service = TSDExportService(db)
    preparation = service.prepare(ee_org['id'], approved['id'], actor=ACTOR)
    workpaper = service.get(ee_org['id'], preparation['id'])
    for field in TSD_AMOUNT_FIELDS:
        assert sum((Decimal(record[field]) for record in workpaper['records']), Decimal('0')) == Decimal(workpaper['totals'][field])
    from ledger import LedgerService
    gl = LedgerService(db).profit_and_loss(ee_org['id'], start='2026-10-01', end='2026-10-31')
    assert Decimal(str(gl['expenses'])) == Decimal(workpaper['totals']['employer_cost'])
    service.review(ee_org['id'], preparation['id'], actor=ACTOR)
    payload = service.export_legacy_annex1_csv(ee_org['id'], preparation['id'], actor=ACTOR)
    assert payload.decode('utf-8') == (
        '1000;1010;1020;1030;1040;1050;1060;1070;1080;1090;1100;1110;1120;1130;1140;1150;1160;1170\r\n'
        '49001010018;Mixed CSV recipient;10;2200.00;;;2200.00;0.00;0.00;0.00;726.00;44.00;2200.00;35.20;17.60;610;700.00;312.58\r\n'
        '49001010018;Mixed CSV recipient;17;400.00;;;400.00;0.00;0.00;0.00;132.00;8.00;400.00;6.40;3.20;;0.00;84.83\r\n'
        '49001010018;Mixed CSV recipient;21;300.00;;;300.00;0.00;0.00;0.00;99.00;6.00;0.00;0.00;0.00;;0.00;64.68\r\n'
        '49001010018;Mixed CSV recipient;24;100.00;;;0.00;0.00;0.00;0.00;0.00;0.00;0.00;0.00;0.00;;0.00;22.00\r\n'
    )


def test_legacy_annex1_csv_refuses_payment_month_after_transition(db, ee_org):
    payroll = PayrollService(db)
    payroll.save_employee(ee_org['id'], actor=ACTOR, name='Future recipient', personal_id='49001010018', gross_salary='2000')
    run = payroll.create_run(ee_org['id'], period='2026-10', actor=ACTOR)
    approved = payroll.approve(ee_org['id'], run['id'], actor=ACTOR)
    with db.transaction() as tx:
        tx.execute("UPDATE pay_runs SET settlement_date='2028-01-05' WHERE organisation_id=? AND id=?", (ee_org['id'], approved['id']))
    service = TSDExportService(db)
    preparation = service.prepare(ee_org['id'], approved['id'], actor=ACTOR)
    service.review(ee_org['id'], preparation['id'], actor=ACTOR)
    with pytest.raises(ValueError, match='December 2027'):
        service.export_legacy_annex1_csv(ee_org['id'], preparation['id'], actor=ACTOR)


def test_workpaper_xlsx_has_three_views(db, ee_org):
    from openpyxl import load_workbook
    service, _, preparation = _reviewed_preparation(db, ee_org)
    book = load_workbook(io.BytesIO(service.workpaper_xlsx(ee_org['id'], preparation['id'])), data_only=True)
    assert book.sheetnames == ['Summary', 'Payments', 'By payment type', 'By recipient']
    assert book['Summary']['B2'].value == len(preparation['records'])
    assert book['Summary']['A8'].value == 'Total payable'
    assert book['Payments'].max_row == 3


def test_tsd_scope_confirmation_and_tenant_gates(db, ee_org, uk_org):
    service = TSDExportService(db)
    with pytest.raises(ValueError, match="Estonian"):
        service.prepare(uk_org["id"], "missing", actor=ACTOR)
    payroll = PayrollService(db)
    payroll.save_employee(
        ee_org["id"], actor=ACTOR, name="No identifier", gross_salary="2000"
    )
    draft = payroll.create_run(ee_org["id"], period="2026-10", actor=ACTOR)
    with pytest.raises(ValueError, match="approved"):
        service.prepare(ee_org["id"], draft["id"], actor=ACTOR)
    approved = payroll.approve(ee_org["id"], draft["id"], actor=ACTOR)
    with pytest.raises(ValueError, match="personal ID"):
        service.prepare(ee_org["id"], approved["id"], actor=ACTOR)


def test_tsd_api_review_gate_and_file_downloads(db, ee_org, monkeypatch):
    run = _approved_run(db, ee_org)
    monkeypatch.setenv("FASTACCOUNTS_DB", db.path)
    monkeypatch.setenv("DB_URL", "")
    from api_app import api

    client = TestClient(api)
    headers = {"X-Test-User": ACTOR}
    base = f"/organisations/{ee_org['id']}/pay-runs/{run['id']}/tsd"
    prepared = client.post(base + "/prepare", headers=headers)
    assert prepared.status_code == 200, prepared.text
    preparation_id = prepared.json()["id"]
    assert client.get(base + f"/{preparation_id}/export.xbrl.xml", headers=headers).status_code == 422
    reviewed = client.post(base + f"/{preparation_id}/review", headers=headers)
    assert reviewed.status_code == 200, reviewed.text
    xml = client.get(base + f"/{preparation_id}/export.xbrl.xml", headers=headers)
    csv_response = client.get(base + f"/{preparation_id}/export.csv", headers=headers)
    legacy = client.get(base + f"/{preparation_id}/export-annex1.csv", headers=headers)
    workpaper = client.get(base + f"/{preparation_id}/workpaper.xlsx", headers=headers)
    assert xml.status_code == 200 and xml.content.startswith(b"<?xml")
    assert xml.headers["content-disposition"].endswith('tsd-2026-10.xbrl.xml"')
    assert csv_response.status_code == 200
    assert csv_response.content.startswith(b"\xef\xbb\xbfdeclaration,")
    assert legacy.status_code == 200 and legacy.content.startswith(b'1000;1010;1020;')
    assert workpaper.status_code == 200 and workpaper.content.startswith(b'PK')
    assert client.get(base + f'/{preparation_id}/workpaper.xlsx', headers={'X-Test-User': 'intruder@example.test'}).status_code == 403


def test_tsd_provider_is_export_only():
    with pytest.raises(PermissionError, match="direct TSD filing is not enabled"):
        EMTATSDExportProvider().submit({"declaration": "TSD"})


def test_services_contract_income_line_can_prepare_and_export(db, ee_org):
    payroll = PayrollService(db)
    payroll.save_employee(ee_org["id"], actor=ACTOR, name="Contractor Test",
                          personal_id="49001010018", pay_basis="income_line")
    run = payroll.create_run(ee_org["id"], period="2026-10", actor=ACTOR)
    kind = next(t for t in payroll.component_types(ee_org["id"]) if t["preset_key"] == "contract")
    payroll.update_item(ee_org["id"], run["id"], run["items"][0]["id"], actor=ACTOR,
                        components=[{"type_id": kind["id"], "amount": "2000"}])
    payroll.approve(ee_org["id"], run["id"], actor=ACTOR)
    service = TSDExportService(db)
    prepared = service.prepare(ee_org["id"], run["id"], actor=ACTOR)
    assert prepared["records"][0]["payment_classification"] == "CONTRACTOR_FEE"
    service.review(ee_org["id"], prepared["id"], actor=ACTOR)
    rows = list(csv.DictReader(io.StringIO(service.export_csv(
        ee_org["id"], prepared["id"], actor=ACTOR).decode("utf-8-sig"))))
    assert rows[0]["pay_basis"] == "income_line"
    assert rows[0]["gross_amount"] == "2000.00"
    root = ET.fromstring(service.export_xbrl_gl(ee_org["id"], prepared["id"], actor=ACTOR))
    assert root.findtext(".//{*}payment_classification") == "CONTRACTOR_FEE"


def test_tsd_payment_date_uses_salary_payout_and_shifts_weekend_deadline(db, ee_org):
    PaydaySettingsService(db).save(ee_org['id'], start_date='2026-01-01', rule_start_date='2026-01-01',
                                   rule='next_month_day', day=10, actor=ACTOR)
    run = _approved_run(db, ee_org)
    PayoutService(db).create(ee_org['id'], run['id'], actor=ACTOR, payout_date='2026-11-05')
    preparation = TSDExportService(db).prepare(ee_org['id'], run['id'], actor=ACTOR)
    assert preparation['payment_date'] == '2026-11-05'
    assert preparation['filing_deadline'] == '2026-12-10'


def test_approved_tsd_lines_group_codes_and_reconcile_frozen_levies(db, ee_org):
    payroll = PayrollService(db)
    employee = payroll.save_employee(ee_org['id'], actor=ACTOR, name='Mixed payment recipient',
                                     personal_id='49001010018', gross_salary='2000',
                                     apply_tax_free_minimum=True)
    run = payroll.create_run(ee_org['id'], period='2026-10', actor=ACTOR)
    types = {row['preset_key']: row for row in payroll.component_types(ee_org['id'])}
    types['board_fee'] = payroll.save_component_type(ee_org['id'], actor=ACTOR,
        name='Board member fee', direction='income', tsd_payment_type=21,
        income_tax=True, social_tax=True, ui_employee=False, ui_employer=False,
        funded_pension=True, social_minimum=False)
    payroll.update_item(ee_org['id'], run['id'], run['items'][0]['id'], actor=ACTOR, components=[
        {'type_id': types['bonus']['id'], 'amount': '100'},
        {'type_id': types['board_fee']['id'], 'amount': '300'},
        {'type_id': types['contract']['id'], 'amount': '400'},
        {'type_id': types['employer_sick_pay']['id'], 'amount': '100'},
    ])
    approved = payroll.approve(ee_org['id'], run['id'], actor=ACTOR)
    item = approved['items'][0]
    lines = db.rows("SELECT * FROM pay_item_tsd_lines WHERE organisation_id=? AND source_kind='pay_run_item' AND source_id=? ORDER BY payment_type",
                    (ee_org['id'], item['id']))
    assert [line['payment_type'] for line in lines] == [10, 17, 21, 24]
    assert Decimal(str(next(line for line in lines if line['payment_type'] == 10)['amount'])) == Decimal('2100.00')
    assert sum((Decimal(line['amount']) for line in lines), Decimal(0)) == Decimal(item['gross'])
    for field, item_field in [('social_tax', 'social_tax'), ('ui_employee', 'ui_employee'),
                              ('ui_employer', 'ui_employer'), ('ii_pillar', 'funded_pension'),
                              ('income_tax', 'income_tax'), ('exemption_610', 'tax_free_minimum')]:
        assert sum((Decimal(str(line[field])) for line in lines), Decimal(0)) == Decimal(str(item[item_field]))
    assert sum((Decimal(str(line['social_tax_base'])) for line in lines), Decimal(0)) == Decimal(str(item['social_tax_base']))
    assert sum((Decimal(str(line['ui_employee_base'])) for line in lines), Decimal(0)) == Decimal('2500.00')
    assert sum((Decimal(str(line['ui_employer_base'])) for line in lines), Decimal(0)) == Decimal('2500.00')
    assert sum((Decimal(str(line['ii_pillar_base'])) for line in lines), Decimal(0)) == Decimal('2800.00')
    assert Decimal(str(next(line for line in lines if line['payment_type'] == 24)['social_tax_base'])) == 0
    assert Decimal(str(next(line for line in lines if line['payment_type'] == 10)['exemption_610'])) == Decimal(str(item['tax_free_minimum']))


def test_tsd_legacy_missing_context_warns_without_mutating_frozen_values(db, ee_org):
    payroll = PayrollService(db)
    payroll.save_employee(ee_org['id'], actor=ACTOR, name='Legacy recipient',
                          personal_id='49001010018', gross_salary='2000')
    draft = payroll.create_run(ee_org['id'], period='2026-10', actor=ACTOR)
    item_id = draft['items'][0]['id']
    with db.transaction() as tx:
        tx.execute('DELETE FROM pay_item_context WHERE organisation_id=? AND item_id=?',
                   (ee_org['id'], item_id))
        tx.execute("UPDATE pay_runs SET status='Approved' WHERE organisation_id=? AND id=?",
                   (ee_org['id'], draft['id']))
    frozen_gross = db.scalar('SELECT gross FROM pay_run_items WHERE organisation_id=? AND id=?',
                             (ee_org['id'], item_id))
    prepared = TSDExportService(db).prepare(ee_org['id'], draft['id'], actor=ACTOR)
    assert 'Legacy run: per-type split unavailable' in prepared['warnings']
    assert len(prepared['records']) == 1 and prepared['records'][0]['tsd_payment_type'] == 10
    assert db.scalar('SELECT gross FROM pay_run_items WHERE organisation_id=? AND id=?',
                     (ee_org['id'], item_id)) == frozen_gross


def test_tsd_breakdown_assigns_minimum_topup_and_exemption_to_salary_first():
    from payroll_tsd_breakdown import calculate_tsd_breakdown

    all_flags = {'income_tax': 1, 'social_tax': 1, 'ui_employee': 1,
                 'ui_employer': 1, 'funded_pension': 1}
    item = {'gross': '1000.00', 'pay_basis': 'monthly', 'tsd_payment_type': 10,
            'social_tax_base': '1200.00', 'social_tax': '396.00',
            'ui_employee': '16.00', 'ui_employer': '8.00', 'funded_pension': '20.00',
            'tax_free_minimum': '700.00', 'income_tax': '62.39'}
    component = {'amount': '400.00', 'direction': 'income', 'tsd_payment_type': 17, **all_flags}
    lines = calculate_tsd_breakdown(item=item, components=[component], context={})
    by_type = {line['payment_type']: line for line in lines}
    assert by_type[10]['amount'] == Decimal('600.00')
    assert by_type[17]['amount'] == Decimal('400.00')
    assert by_type[10]['social_tax_minimum_topup'] == Decimal('200.00')
    assert by_type[17]['social_tax_minimum_topup'] == Decimal('0.00')
    assert by_type[10]['exemption_610'] == Decimal('578.40')
    assert by_type[17]['exemption_610'] == Decimal('121.60')
    for field, expected in [('social_tax', '396.00'), ('ui_employee', '16.00'),
                            ('ui_employer', '8.00'), ('ii_pillar', '20.00'),
                            ('exemption_610', '700.00'), ('income_tax', '62.39')]:
        assert sum((line[field] for line in lines), Decimal(0)) == Decimal(expected)
    assert sum((line['social_tax_base'] for line in lines), Decimal(0)) == Decimal('1200.00')
    assert sum((line['ui_employee_base'] for line in lines), Decimal(0)) == Decimal('1000.00')
    assert sum((line['ui_employer_base'] for line in lines), Decimal(0)) == Decimal('1000.00')
    assert sum((line['ii_pillar_base'] for line in lines), Decimal(0)) == Decimal('1000.00')
