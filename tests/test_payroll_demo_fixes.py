"""Offline bureau payroll lifecycle, component snapshots and journal contracts."""
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from core_utils import money
from payroll import PayrollConflict, PayrollService, calculate_pay, declaration_due_date
from payroll_components import FLAGS

ACTOR = 'owner@example.test'


def test_full_catalogue_defaults_and_descriptions(db, ee_org, uk_org):
    from payroll_components import PRESETS, PRESET_DESCRIPTIONS, seed_types
    service = PayrollService(db)
    types = {t['preset_key']: t for t in service.component_types(ee_org['id'])}
    assert len(types) == len(PRESETS) == 27
    assert set(types) == set(PRESET_DESCRIPTIONS)
    for key, name, direction, flags, role in PRESETS:
        actual = types[key]
        assert actual['name'] == name and actual['description'] == PRESET_DESCRIPTIONS[key]
        assert actual['direction'] == direction and actual['liability_role'] == role
        assert tuple(actual[f] for f in FLAGS) == flags
        if direction == 'deduction':
            assert flags == (0, 0, 0, 0, 0, 0)
            assert role == ('PAYROLL_BAILIFF' if key == 'bailiff' else 'PAYROLL_DEDUCTION')
    assert tuple(types['employer_sick_pay'][f] for f in FLAGS) == (1, 0, 0, 0, 0, 0)
    assert not any(types['expense_reimbursement'][f] for f in FLAGS)
    assert all(types[key][f] for key in ('severance', 'author', 'parental') for f in FLAGS)
    with db.transaction() as tx:
        seed_types(tx, ee_org['id'])
    assert len(service.component_types(ee_org['id'])) == 27
    assert db.scalar('SELECT COUNT(*) FROM pay_component_types WHERE organisation_id=?', (uk_org['id'],)) == 0


@pytest.mark.parametrize('key', ['salary_advance', 'advance_recovery', 'alimony', 'union_fee', 'agreed_deduction'])
def test_new_net_deductions_post_to_other_liability(db, ee_org, key):
    service, person, run, types = setup_run(db, ee_org)
    before = run['items'][0]
    run = edit(service, ee_org, run, [dict(type_id=types[key]['id'], amount='25')])
    item = run['items'][0]
    assert money(item['net']) == money(before['net']) - Decimal('25')
    for field in ('gross', 'income_tax', 'social_tax', 'funded_pension', 'ui_employee', 'ui_employer', 'employer_cost'):
        assert item[field] == before[field]
    approved = service.approve(ee_org['id'], run['id'], actor=ACTOR)
    lines = balanced(db, approved['posted_batch_id'])
    account = db.scalar("SELECT id FROM accounts WHERE organisation_id=? AND system_role='PAYROLL_DEDUCTION'", (ee_org['id'],))
    assert sum(money(line['credit']) for line in lines if line['account_id'] == account) == Decimal('25')


def test_sick_pay_and_documented_reimbursement_use_separate_tax_bases(db, ee_org):
    service, person, run, types = setup_run(db, ee_org)
    before = run['items'][0]
    run = edit(service, ee_org, run, [dict(type_id=types[key]['id'], amount='100')
                                    for key in ('employer_sick_pay', 'expense_reimbursement')])
    item = run['items'][0]
    assert money(item['gross']) == money(before['gross']) + Decimal('200')
    assert money(item['income_tax']) == money(before['income_tax']) + Decimal('22')
    assert money(item['net']) == money(before['net']) + Decimal('178')
    for field in ('social_tax', 'funded_pension', 'ui_employee', 'ui_employer'):
        assert item[field] == before[field]
    balanced(db, service.approve(ee_org['id'], run['id'], actor=ACTOR)['posted_batch_id'])


def setup_run(db, org, **employee):
    service = PayrollService(db)
    person = service.save_employee(org['id'], actor=ACTOR, **(dict(name='Synthetic payroll', gross_salary='2000', personal_id='49001010001') | employee))
    run = service.create_run(org['id'], period='2026-10', actor=ACTOR,
                             hours={person['id']: '160'} if person['pay_basis'] == 'hourly' else None)
    types = {t['preset_key']: t for t in service.component_types(org['id'])}
    return service, person, run, types


def edit(service, org, run, components, **changes):
    return service.update_item(org['id'], run['id'], run['items'][0]['id'], actor=ACTOR, components=components, **changes)


def balanced(db, batch):
    lines = db.rows('SELECT * FROM gl_entries WHERE posting_batch_id=?', (batch,))
    assert lines
    assert sum((money(line['debit']) for line in lines), Decimal(0)) == sum((money(line['credit']) for line in lines), Decimal(0))
    return lines


def test_archive_unarchive_delete_history_and_audit(db, ee_org):
    service, person, run, types = setup_run(db, ee_org)
    before = run['items']
    service.save_employee(ee_org['id'], employee_id=person['id'], active=False, actor=ACTOR)
    assert service.runs(ee_org['id'])[0]['items'] == before
    with pytest.raises(ValueError, match='active employee'):
        service.create_run(ee_org['id'], period='2026-11', actor=ACTOR)
    with pytest.raises(PayrollConflict, match='Archive the employee instead'):
        service.delete_employee(ee_org['id'], person['id'], actor=ACTOR)
    service.save_employee(ee_org['id'], employee_id=person['id'], active=True, actor=ACTOR)
    unused = service.save_employee(ee_org['id'], name='Unused', gross_salary='1000', actor=ACTOR)
    service.delete_employee(ee_org['id'], unused['id'], actor=ACTOR)
    assert [p['id'] for p in service.employees(ee_org['id'])] == [person['id']]
    actions = {row['action'] for row in db.rows('SELECT action FROM audit_events WHERE organisation_id=?', (ee_org['id'],))}
    assert {'payroll.employee.archived', 'payroll.employee.unarchived', 'payroll.employee.deleted'} <= actions


def test_components_rounding_snapshots_exports_and_approval(db, ee_org):
    service, person, run, types = setup_run(db, ee_org, apply_tax_free_minimum=True)
    # Explicit reviewed legacy flags still freeze independently of conservative presets.
    run = edit(service, ee_org, run, [dict(type_id=types['author']['id'], amount='100', **dict(zip(FLAGS, (1, 1, 0, 0, 0, 1)))),
                                    dict(type_id=types['parental']['id'], amount='50', **dict(zip(FLAGS, (1, 0, 0, 0, 0, 0)))),
                                    dict(type_id=types['bailiff']['id'], amount='30')])
    item = run['items'][0]
    assert (item['gross'], item['funded_pension'], item['ui_employee'], item['income_tax'], item['net'], item['social_tax'], item['ui_employer']) == (
        '2150.00', '40.00', '32.00', '303.16', '1744.84', '693.00', '16.00')
    assert item['employer_cost'] == run['total_employer_cost'] == '2859.00'
    assert item['withholding_total'] == '405.16'
    service.save_component_type(ee_org['id'], type_id=types['author']['id'], name='Changed default', income_tax=False, actor=ACTOR)
    service.save_employee(ee_org['id'], employee_id=person['id'], gross_salary='4000', active=False, actor=ACTOR)
    run = edit(service, ee_org, run, [dict(type_id=line['type_id'], amount=line['amount'], name=line['name'],
                                        direction=line['direction'], **{key: line[key] for key in FLAGS}) for line in item['components']])
    assert run['items'][0]['gross'] == '2150.00'
    assert run['items'][0]['components'][0]['name'] == 'Autori tasu'
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(service.workbook_xlsx(ee_org['id'], run['id'])))
    assert book['Palga lisaread'].cell(2, 2).value == 'Autori tasu'
    assert book['Palga lisaread'].cell(4, 4).value == 30
    from pypdf import PdfReader
    text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(service.payslips_pdf(ee_org['id'], run['id']))).pages)
    assert 'Autori tasu' in text and 'Kinnipidamine kohtutäiturile' in text and '30.00' in text
    approved = service.approve(ee_org['id'], run['id'], actor=ACTOR)
    balanced(db, approved['posted_batch_id'])
    for action in [lambda: edit(service, ee_org, run, []), lambda: service.delete_run(ee_org['id'], run['id'], actor=ACTOR)]:
        with pytest.raises(PayrollConflict):
            action()
    with pytest.raises(Exception, match='immutable'):
        with db.transaction() as tx:
            tx.execute('UPDATE pay_item_components SET amount=1 WHERE organisation_id=?', (ee_org['id'],))


def test_hourly_edit_retains_frozen_allowance_and_dates(db, ee_org):
    service, person, run, types = setup_run(db, ee_org, pay_basis='hourly', hourly_rate='8',
        apply_tax_free_minimum=True, employment_start_date='2026-10-16')
    service.save_employee(ee_org['id'], employee_id=person['id'], hourly_rate='12',
        apply_tax_free_minimum=False, employment_start_date=None, actor=ACTOR)
    run = edit(service, ee_org, run, [], hours='50')
    item = run['items'][0]
    assert item['gross'] == '400.00' and item['income_tax'] == '0.00'
    assert item['social_tax_base'] == '457.29' and item['social_tax'] == '150.91'
    assert run['total_employer_cost'] == '554.11'
    with pytest.raises(ValueError):
        edit(service, ee_org, run, [], hours='0')
    assert service.runs(ee_org['id'])[0] == run


def test_independent_flags_minimum_topup_and_net_deduction():
    base = dict(gross_salary='500', funded_pension_percent=2, apply_tax_free_minimum=False)
    flags = {key: False for key in FLAGS}
    author = dict(name='Author', amount='100', direction='income', **(flags | dict(social_minimum=True)))
    result = calculate_pay(base, period='2026-10', components=[author])
    assert result['social_tax'] == Decimal('259.38')  # 165 actual + 94.38 top-up.
    assert result['social_tax_minimum_topup'] == Decimal('94.38')
    assert result['income_tax'] == Decimal('106.04') and result['gross'] == Decimal('600')
    exempt = calculate_pay(base | dict(social_tax_minimum_exemption='student'), period='2026-10', components=[author])
    assert exempt['social_tax_minimum_topup'] == 0 and exempt['social_tax'] == Decimal('165')
    for key, amount_key, delta in [('income_tax', 'income_tax', '22'), ('social_tax', 'social_tax', '33'),
                                   ('ui_employee', 'ui_employee', '1.60'), ('ui_employer', 'ui_employer', '.80'),
                                   ('funded_pension', 'funded_pension', '2')]:
        result = calculate_pay(base | dict(social_tax_minimum_exemption='student'), period='2026-10', components=[author | {key: True}])
        assert result[amount_key] - exempt[amount_key] == Decimal(delta)


def test_custom_type_negative_adjustment_validation_and_atomicity(db, ee_org):
    service, person, run, types = setup_run(db, ee_org)
    custom = service.save_component_type(ee_org['id'], name='Synthetic correction', direction='income',
        actor=ACTOR, **{key: True for key in FLAGS})
    run = edit(service, ee_org, run, [dict(type_id=custom['id'], amount='-100')])
    assert run['total_gross'] == '1900.00'
    before = run
    for amount in ['NaN', 'Infinity', '0', '10.001', '-2000']:
        with pytest.raises(ValueError):
            edit(service, ee_org, run, [dict(type_id=custom['id'], amount=amount)])
        assert service.runs(ee_org['id'])[0] == before
    with pytest.raises(ValueError, match='net pay'):
        edit(service, ee_org, run, [dict(type_id=types['bailiff']['id'], amount='99999')])
    assert service.runs(ee_org['id'])[0] == before


def test_month_end_shape_mappings_and_explicit_settlement(db, ee_org):
    from tsd import TSDExportService
    service, person, run, types = setup_run(db, ee_org)
    run = edit(service, ee_org, run, [dict(type_id=types['bailiff']['id'], amount='50')])
    # Map the bureau's salary payable role onto its original 2210 chart account.
    service.save_posting_mapping(ee_org['id'], 'PAYROLL_SALARY_PAYABLE', '2210', actor=ACTOR)
    run = service.approve(ee_org['id'], run['id'], actor=ACTOR)
    salary_account = db.scalar("SELECT id FROM accounts WHERE organisation_id=? AND code='2210'", (ee_org['id'],))
    salary_lines = [l for l in balanced(db, run['posted_batch_id']) if l['account_id'] == salary_account]
    assert [(money(l['debit']), money(l['credit'])) for l in salary_lines] == [(Decimal(0), Decimal('2000')), (Decimal('546.16'), Decimal(0))]
    assert run['declaration_state'] == 'none' and run['declaration_due_date'] == '2026-11-10'
    assert declaration_due_date('2026-12').isoformat() == '2027-01-11'  # 10 January is Sunday, so EMTA moves the due date to Monday.
    TSDExportService(db).prepare(ee_org['id'], run['id'], actor=ACTOR)
    assert service.runs(ee_org['id'])[0]['declaration_state'] == 'prepared'
    before = db.rows('SELECT * FROM gl_entries WHERE posting_batch_id=? ORDER BY id', (run['posted_batch_id'],))
    # Changing the accrued mapping after approval must not strand the old balance.
    old_accrual = db.scalar("SELECT id FROM accounts WHERE organisation_id=? AND system_role='PAYROLL_SOCIAL_ACCRUED'", (ee_org['id'],))
    service.save_posting_mapping(ee_org['id'], 'PAYROLL_SOCIAL_ACCRUED', '2240', actor=ACTOR)
    settled = service.settle_declaration(ee_org['id'], run['id'], actor=ACTOR)
    lines = balanced(db, settled['declaration_batch_id'])
    assert any(l['account_id'] == old_accrual and money(l['debit']) == Decimal('660') for l in lines)
    assert settled['declaration_state'] == 'settled' and not settled['declaration_overdue']
    assert db.scalar('SELECT posting_date FROM posting_batches WHERE id=?', (settled['declaration_batch_id'],)) == '2026-11-10'
    with pytest.raises(PayrollConflict):
        service.settle_declaration(ee_org['id'], run['id'], actor=ACTOR)
    assert db.rows('SELECT * FROM gl_entries WHERE posting_batch_id=? ORDER BY id', (run['posted_batch_id'],)) == before


def test_new_endpoints_access_tenancy_and_validation(db, ee_org, monkeypatch):
    from api_app import api
    from organisations import OrganisationService
    monkeypatch.setenv('FASTACCOUNTS_DB', db.path)
    monkeypatch.setenv('DB_URL', '')
    service, person, run, types = setup_run(db, ee_org)
    other = OrganisationService(db).create(name='Synthetic other payroll', country_code='EE', entity_type='EE_OU', owner_email=ACTOR)
    client = TestClient(api)
    root = f"/organisations/{ee_org['id']}"
    headers = {'X-Test-User': ACTOR}
    item_url = root + f"/pay-runs/{run['id']}/items/{run['items'][0]['id']}"
    assert client.patch(item_url, json={'components': []}).status_code == 401
    assert client.patch(item_url, headers={'X-Test-User': 'intruder@example.test'}, json={'components': []}).status_code == 403
    assert client.delete(root + f"/employees/{person['id']}", headers=headers).status_code == 409
    assert client.patch(item_url, headers=headers, json={'components': [{'type_id': types['author']['id'], 'amount': '100'}]}).status_code == 200
    assert client.patch(item_url, headers=headers, json={'components': None}).status_code == 422
    foreign_type = service.component_types(other['id'])[0]
    assert client.patch(item_url, headers=headers, json={'components': [{'type_id': foreign_type['id'], 'amount':'100'}]}).status_code == 404
    with pytest.raises(KeyError):
        service.delete_employee(other['id'], person['id'], actor=ACTOR)
    with pytest.raises(KeyError):
        service.update_item(other['id'], run['id'], run['items'][0]['id'], components=[], actor=ACTOR)
    assert client.get(root + '/payroll/posting-mappings', headers=headers).status_code == 200
    created = client.post(root + '/payroll/component-types', headers=headers, json={'name':'Custom', 'income_tax':True})
    assert created.status_code == 201
    assert client.put(root + '/payroll/component-types/' + created.json()['id'], headers=headers,
        json={'name':'Changed', 'social_tax':True}).status_code == 200


def test_migration_0034_is_rerun_safe_and_preserves_posted_data(db, ee_org):
    service, person, run, types = setup_run(db, ee_org)
    run = service.approve(ee_org['id'], run['id'], actor=ACTOR)
    before = service.runs(ee_org['id'])
    entries = db.rows('SELECT * FROM gl_entries ORDER BY id')
    with db.transaction() as tx:
        tx.execute("DELETE FROM schema_migrations WHERE version='0034_payroll_demo_fixes'")
    assert db.migrate() == ['0034_payroll_demo_fixes']
    assert service.runs(ee_org['id']) == before
    assert db.rows('SELECT * FROM gl_entries ORDER BY id') == entries
    assert db.rows('PRAGMA foreign_key_check') == []


def test_settlement_lock_rolls_back_and_concurrent_requests_post_once(db, ee_org):
    from concurrent.futures import ThreadPoolExecutor
    from tsd import TSDExportService
    service, person, run, types = setup_run(db, ee_org)
    run = service.approve(ee_org['id'], run['id'], actor=ACTOR)
    TSDExportService(db).prepare(ee_org['id'], run['id'], actor=ACTOR)
    with db.transaction() as tx:
        tx.execute("UPDATE fiscal_periods SET status='Locked' WHERE organisation_id=? AND code='2026-11'", (ee_org['id'],))
    with pytest.raises(ValueError, match='locked'):
        service.settle_declaration(ee_org['id'], run['id'], actor=ACTOR)
    assert service.runs(ee_org['id'])[0]['declaration_state'] == 'prepared'
    assert db.scalar('SELECT COUNT(*) FROM pay_run_settlements') == 0
    with db.transaction() as tx:
        tx.execute("UPDATE fiscal_periods SET status='Open' WHERE organisation_id=? AND code='2026-11'", (ee_org['id'],))
    def settle():
        try:
            service.settle_declaration(ee_org['id'], run['id'], actor=ACTOR)
            return 'settled'
        except PayrollConflict:
            return 'conflict'
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(lambda _: settle(), range(2))) == ['conflict', 'settled']
    assert db.scalar("SELECT COUNT(*) FROM posting_batches WHERE voucher_type='payroll_tsd'") == 1


def test_components_in_reviewed_push_and_tsd_taxable_base(db, ee_org):
    import json
    from integration_service import AccountMappingService
    from push_service import PushService
    from tsd import TSDExportService
    service, person, run, types = setup_run(db, ee_org)
    run = edit(service, ee_org, run, [dict(type_id=types['parental']['id'], amount='100', **{f: False for f in FLAGS}),
                                    dict(type_id=types['bailiff']['id'], amount='10')])
    run = service.approve(ee_org['id'], run['id'], actor=ACTOR)
    for row in db.rows('SELECT DISTINCT a.code FROM gl_entries g JOIN accounts a ON a.id=g.account_id AND a.organisation_id=g.organisation_id WHERE g.posting_batch_id=?', (run['posted_batch_id'],)):
        AccountMappingService(db).create(ee_org['id'], 'merit', local_account_code=row['code'], external_account_code=row['code'], status='Confirmed', actor=ACTOR)
    push = PushService(db).prepare(ee_org['id'], 'merit', 'pay_runs', run['id'], ACTOR)
    payload = json.loads(push['payload_json'])
    assert payload['items'][0]['components'] == run['items'][0]['components']
    assert sum((money(line['debit']) for line in payload['lines']), Decimal(0)) == sum((money(line['credit']) for line in payload['lines']), Decimal(0))
    PushService(db).review(ee_org['id'], push['id'], ACTOR)
    workpaper = TSDExportService(db).prepare(ee_org['id'], run['id'], actor=ACTOR)
    assert workpaper['records'][0]['income_tax_base'] == '1928.00'
    assert workpaper['records'][0]['income_tax'] == '424.16'


def test_declaration_overdue_boundary_and_foreign_mappings(db, ee_org, monkeypatch):
    import payroll
    from datetime import date
    from organisations import OrganisationService
    service, person, run, types = setup_run(db, ee_org)
    service.approve(ee_org['id'], run['id'], actor=ACTOR)
    class DueDay(date):
        @classmethod
        def today(cls):
            return cls(2026, 11, 10)
    monkeypatch.setattr(payroll, 'date', DueDay)
    assert not service.runs(ee_org['id'])[0]['declaration_overdue']
    class LateDay(date):
        @classmethod
        def today(cls):
            return cls(2026, 11, 11)
    monkeypatch.setattr(payroll, 'date', LateDay)
    assert service.runs(ee_org['id'])[0]['declaration_overdue']
    other = OrganisationService(db).create(name='Other chart', country_code='EE', entity_type='EE_OU', owner_email=ACTOR)
    # A chart code that belongs only to another organisation is never selectable.
    with db.transaction() as tx:
        tx.execute("UPDATE accounts SET code='SYNTHETIC-OTHER' WHERE organisation_id=? AND code='2251'", (other['id'],))
    with pytest.raises(ValueError, match='active account'):
        service.save_posting_mapping(ee_org['id'], 'PAYROLL_BAILIFF', 'SYNTHETIC-OTHER', actor=ACTOR)


def test_legacy_minimum_wage_and_mapping_account_type_are_revalidated(db, ee_org):
    service, person, run, types = setup_run(db, ee_org, gross_salary='1000')
    # An older draft may not have a stored minimum_wage, but edits must still
    # enforce the period's minimum rather than treating a missing value as zero.
    with db.transaction() as tx:
        tx.execute('UPDATE pay_run_items SET minimum_wage=NULL WHERE id=?', (run['items'][0]['id'],))
    correction = service.save_component_type(ee_org['id'], name='Gross correction', actor=ACTOR,
        **{key: True for key in FLAGS})
    with pytest.raises(ValueError, match='minimum wage'):
        edit(service, ee_org, run, [dict(type_id=correction['id'], amount='-100')])
    service.save_posting_mapping(ee_org['id'], 'PAYROLL_SOCIAL_ACCRUED', '2253', actor=ACTOR)
    with db.transaction() as tx:
        tx.execute("UPDATE accounts SET account_type='Expense' WHERE organisation_id=? AND code='2253'", (ee_org['id'],))
    with pytest.raises(ValueError, match='required payroll account type'):
        service.approve(ee_org['id'], run['id'], actor=ACTOR)
    assert service.runs(ee_org['id'])[0]['status'] == 'Draft'


def test_new_payroll_copy_is_present_in_all_catalogues():
    from web import i18n
    from payroll_components import FLAG_LABELS
    keys = ['Archive', 'Unarchive', 'Archived employees', 'Delete employee',
            'This employee has been included in a pay run. Archive the employee instead.',
            'Payment types', 'Payment components', 'Edit draft pay', 'Payroll accounts',
            'Month-end declarations', 'Overdue payroll declarations', 'Reclassify TSD liabilities',
            'Kinnipidamine kohtutäiturile', 'Autori tasu', 'Vanemapuhkused',
            *FLAG_LABELS]
    for language in ('en', 'et', 'lv', 'lt'):
        catalogue = i18n.catalog(language)['workspace']
        assert all(catalogue.get(key) for key in keys)


def test_zero_liability_declaration_does_not_fabricate_zero_postings(db, ee_org):
    from tsd import TSDExportService
    service, person, run, types = setup_run(db, ee_org, gross_salary='1000', social_tax_minimum_exemption='student')
    correction = service.save_component_type(ee_org['id'], name='Replace gross', actor=ACTOR,
        **{key: True for key in FLAGS})
    run = edit(service, ee_org, run, [dict(type_id=correction['id'], amount='-1000'),
        dict(type_id=types['parental']['id'], amount='2000', **{f: False for f in FLAGS})])
    run = service.approve(ee_org['id'], run['id'], actor=ACTOR)
    assert run['total_gross'] == run['total_net'] == run['total_employer_cost'] == '2000.00'
    TSDExportService(db).prepare(ee_org['id'], run['id'], actor=ACTOR)
    settled = service.settle_declaration(ee_org['id'], run['id'], actor=ACTOR)
    assert settled['declaration_state'] == 'settled' and settled['declaration_batch_id'] is None
    assert db.scalar('SELECT COUNT(*) FROM posting_batches WHERE organisation_id=?', (ee_org['id'],)) == 1
    with pytest.raises(PayrollConflict):
        service.settle_declaration(ee_org['id'], run['id'], actor=ACTOR)


def test_custom_payment_names_and_levy_labels_in_exports():
    from documents import payroll_pdf
    from excel_export import pay_run_workbook
    from openpyxl import load_workbook
    from pypdf import PdfReader
    item = dict(employee_name='Synthetic employee', board_member=False, components=[
        dict(name='=Synthetic custom payment', amount='10.00', direction='income', **{key: True for key in FLAGS})])
    organisation = dict(country_code='EE', name='Synthetic organisation')
    run = dict(period='2026-10', status='Draft', items=[item])
    text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(payroll_pdf(organisation, run))).pages)
    assert '=Synthetic custom payment' in text and 'workspace.' not in text
    assert 'Töötaja töötuskindlustusmakse' in text and 'Kogumispension' in text
    book = load_workbook(BytesIO(pay_run_workbook(organisation, run, [item])))
    name = book['Palga lisaread'].cell(2, 2)
    assert name.value == '=Synthetic custom payment' and name.data_type == 's'


def test_explicit_reclassification_does_not_require_in_app_tsd_files(db, ee_org):
    service, person, run, types = setup_run(db, ee_org, personal_id=None)
    run = service.approve(ee_org['id'], run['id'], actor=ACTOR)
    assert run['declaration_state'] == 'none'
    settled = service.settle_declaration(ee_org['id'], run['id'], actor=ACTOR)
    assert settled['declaration_state'] == 'settled'
    balanced(db, settled['declaration_batch_id'])
    assert db.scalar("SELECT COUNT(*) FROM sync_runs WHERE provider='emta'") == 0


def test_hourly_gross_adjustments_preserve_minimum_pay(db, ee_org):
    service, person, run, types = setup_run(db, ee_org, pay_basis='hourly', hourly_rate='10')
    correction = service.save_component_type(ee_org['id'], name='Hourly gross correction', actor=ACTOR,
        **{key: True for key in FLAGS})
    with pytest.raises(ValueError, match='minimum wage'):
        edit(service, ee_org, run, [dict(type_id=correction['id'], amount='-800')])
    assert service.runs(ee_org['id'])[0] == run
