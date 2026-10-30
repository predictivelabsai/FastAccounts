from io import BytesIO

import pytest
from openpyxl import Workbook
import base64
from fastapi.testclient import TestClient

from payroll_dimension_imports import PayrollDimensionImportService
from api_app import api
from core_utils import new_id, utc_now
from database import reset_database_cache


def test_csv_dimension_import_requires_review_and_creates_position_group(db, ee_org):
    service = PayrollDimensionImportService(db)
    batch = service.stage(ee_org['id'], 'position', 'code,name,group_code,group_name\nDEV,Developer,TECH,Technical\n',
                          file_name='positions.csv', actor='owner@example.test')
    assert batch['status'] == 'staged'
    assert batch['total_rows'] == 1 and batch['error_count'] == 0
    assert db.scalar('SELECT COUNT(*) FROM payroll_dimensions WHERE organisation_id=? AND kind IN (\'position\',\'position_group\')',
                     (ee_org['id'],)) == 0
    result = service.commit(ee_org['id'], batch['id'], actor='owner@example.test')
    assert result['status'] == 'committed'
    group = db.one("SELECT * FROM payroll_dimensions WHERE organisation_id=? AND kind='position_group'", (ee_org['id'],))
    position = db.one("SELECT * FROM payroll_dimensions WHERE organisation_id=? AND kind='position'", (ee_org['id'],))
    assert group['code'] == 'TECH' and position['parent_id'] == group['id']


def test_csv_import_stages_errors_duplicates_and_skips_existing(db, ee_org):
    service = PayrollDimensionImportService(db)
    content = 'code,name\nOPS,Operations\nOPS,Duplicate\n,Missing code\n'
    batch = service.stage(ee_org['id'], 'department', content, file_name='departments.csv', actor='owner@example.test')
    assert batch['total_rows'] == 3 and batch['error_count'] == 1 and batch['warning_count'] == 1
    assert db.scalar('SELECT COUNT(*) FROM payroll_dimensions WHERE organisation_id=?', (ee_org['id'],)) == 0
    applied = service.commit(ee_org['id'], batch['id'], actor='owner@example.test')
    assert applied['rows'][0]['entity_id']
    assert applied['rows'][1]['warning'] == 'Existing record skipped'
    assert applied['rows'][2]['error']
    assert db.scalar('SELECT COUNT(*) FROM payroll_dimensions WHERE organisation_id=?', (ee_org['id'],)) == 1
    repeat = service.stage(ee_org['id'], 'department', 'code,name\nOPS,Operations\n', file_name='again.csv', actor='owner@example.test')
    assert repeat['warning_count'] == 1


def test_xlsx_import_creates_tenant_scoped_project(db, ee_org, uk_org):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['code', 'name'])
    sheet.append(['WEB', 'Website project'])
    content = BytesIO()
    workbook.save(content)
    service = PayrollDimensionImportService(db)
    batch = service.stage(ee_org['id'], 'project', content.getvalue(), file_name='projects.xlsx', actor='owner@example.test')
    result = service.commit(ee_org['id'], batch['id'], actor='owner@example.test')
    project = db.one("SELECT * FROM objects WHERE organisation_id=? AND object_type='project'", (ee_org['id'],))
    assert project['code'] == 'WEB' and result['rows'][0]['entity_id'] == project['id']
    with pytest.raises(KeyError):
        service.get(uk_org['id'], batch['id'])
    assert db.scalar("SELECT COUNT(*) FROM objects WHERE organisation_id=? AND object_type='project'", (uk_org['id'],)) == 0


def test_xlsx_formula_and_unsupported_file_are_rejected(db, ee_org):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['code', 'name'])
    sheet.append(['=1+1', 'Formula code'])
    content = BytesIO()
    workbook.save(content)
    service = PayrollDimensionImportService(db)
    batch = service.stage(ee_org['id'], 'cost_centre', content.getvalue(), file_name='centres.xlsx', actor='owner@example.test')
    assert batch['error_count'] == 1
    with pytest.raises(ValueError, match='CSV or XLSX'):
        service.stage(ee_org['id'], 'department', b'bad', file_name='data.json', actor='owner@example.test')


def test_dimension_import_api_reviews_and_applies_csv_and_xlsx(db, ee_org, uk_org, monkeypatch):
    monkeypatch.setenv('FASTACCOUNTS_DB', db.path)
    monkeypatch.setenv('DB_URL', '')
    reset_database_cache()
    with db.transaction() as tx:
        tx.execute('INSERT INTO memberships(id,organisation_id,email,role,created_at) VALUES (?,?,?,?,?)',
                   (new_id(), ee_org['id'], 'viewer@example.test', 'viewer', utc_now()))
    base = f"/organisations/{ee_org['id']}/payroll/dimension-imports"
    with TestClient(api, headers={'X-Test-User': 'owner@example.test'}) as client:
        staged = client.post(base, json={'kind': 'department', 'file_name': 'departments.csv',
                                         'csv_content': 'code,name\nOPS,Operations\n'})
        assert staged.status_code == 201
        batch_id = staged.json()['id']
        viewer = {'X-Test-User': 'viewer@example.test'}
        assert client.get(base + '/' + batch_id, headers=viewer).status_code == 200
        assert client.post(base + '/' + batch_id + '/commit', headers=viewer).status_code == 403
        assert client.get(f"/organisations/{uk_org['id']}/payroll/dimension-imports/{batch_id}").status_code == 404
        assert client.post(base + '/' + batch_id + '/commit').json()['status'] == 'committed'
        workbook = Workbook()
        workbook.active.append(['code', 'name'])
        workbook.active.append(['INT', 'Integration project'])
        binary = BytesIO()
        workbook.save(binary)
        xlsx = client.post(base, json={'kind': 'project', 'file_name': 'projects.xlsx',
                                       'xlsx_base64': base64.b64encode(binary.getvalue()).decode('ascii')})
        assert xlsx.status_code == 201
        assert client.post(base + '/' + xlsx.json()['id'] + '/discard').json()['status'] == 'discarded'
