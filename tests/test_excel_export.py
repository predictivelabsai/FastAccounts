from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from database import reset_database_cache
from excel_export import (employees_workbook, general_ledger_workbook,
                          invoices_workbook, pay_run_workbook,
                          vat_returns_workbook)
from payroll import PayrollService


ACTOR = "owner@example.test"


def workbook(content: bytes):
    return load_workbook(BytesIO(content), data_only=True)


def payroll_fixture():
    organisation = {"id": "org-ee", "name": "Synthetic Eesti OÜ", "country_code": "EE"}
    items = [
        {
            "employee_name": "Ada Example", "personal_id": None, "pay_basis": "monthly",
            "gross": "2000.00", "funded_pension": "40.00", "ui_employee": "32.00",
            "income_tax": "270.16", "social_tax": "660.00", "ui_employer": "16.00",
            "net": "1657.84", "employer_cost": "2676.00",
        },
        {
            "employee_name": "Ben Example", "personal_id": "39901010000", "pay_basis": "board_fee",
            "gross": "500.00", "funded_pension": "0.00", "ui_employee": "0.00",
            "income_tax": "110.00", "social_tax": "165.00", "ui_employer": "0.00",
            "net": "390.00", "employer_cost": "665.00",
        },
    ]
    run = {
        "id": "run-1", "period": "2026-07", "status": "Approved",
        "total_gross": "2500.00", "total_net": "2047.84", "total_employer_cost": "3341.00",
    }
    return organisation, run, items


def test_pay_run_workbook_reconciles_numeric_totals():
    organisation, run, items = payroll_fixture()
    book = workbook(pay_run_workbook(organisation, run, items))
    assert book.sheetnames == ["Summary", "Items"]
    assert [cell.value for cell in book["Items"][1]] == [
        "Name", "Personal ID", "Pay basis", "Gross", "Funded pension",
        "Employee unemployment insurance", "Income tax", "Social tax",
        "Employer unemployment insurance", "Net", "Employer cost",
    ]
    assert book["Items"].max_row == 4
    assert all(cell.font.bold for cell in book["Items"][1])
    totals = book["Items"][4]
    assert totals[0].value == "Total"
    item_totals = {
        "total_gross": Decimal(str(totals[3].value)),
        "total_net": Decimal(str(totals[9].value)),
        "total_employer_cost": Decimal(str(totals[10].value)),
    }
    assert {key: value.quantize(Decimal("0.01")) for key, value in item_totals.items()} == {
        key: Decimal(run[key]).quantize(Decimal("0.01")) for key in item_totals
    }
    assert book["Items"][2][3].number_format == "#,##0.00"
    assert book["Summary"].max_row == 9


def test_pay_run_accounting_sheet_is_optional():
    organisation, run, items = payroll_fixture()
    without = workbook(pay_run_workbook(organisation, run, items))
    assert "Accounting" not in without.sheetnames
    lines = [
        {"account_code": "6110", "account_name": "Salary expense", "memo": None,
         "debit": "2500.00", "credit": None},
        {"account_code": "2210", "account_name": "Payroll payable", "memo": "Payroll 2026-07",
         "debit": None, "credit": "2500.00"},
    ]
    included = workbook(pay_run_workbook(organisation, run, items, lines))
    assert included.sheetnames == ["Summary", "Items", "Accounting"]
    assert included["Accounting"].max_row == 4
    assert Decimal(str(included["Accounting"][4][3].value)) == Decimal("2500")
    assert Decimal(str(included["Accounting"][4][4].value)) == Decimal("2500")


def test_report_workbooks_render_real_fields_and_none_values():
    organisation = {"name": "Synthetic Eesti OÜ", "country_code": "EE"}
    ledger = workbook(general_ledger_workbook(organisation, [{
        "entry_date": None, "voucher_code": "J-1", "voucher_type": "journal",
        "account_code": "1000", "account_name": "Cash", "memo": None,
        "debit": "10.00", "credit": None,
    }]))
    assert ledger["General Ledger"].max_row == 3
    assert ledger["General Ledger"][2][0].value in (None, "")

    invoices = workbook(invoices_workbook(organisation, [{
        "number": None, "contact_name": None, "issue_date": "2026-01-01", "due_date": None,
        "currency": "EUR", "subtotal": None, "tax_total": "0.00", "total": None, "status": "Draft",
    }]))
    assert invoices["Invoices"].max_row == 3
    assert invoices["Invoices"][3][0].value == "Total"

    employees = workbook(employees_workbook(organisation, [{
        "name": "Null Example", "personal_id": None, "email": None, "pay_basis": "hourly",
        "gross_salary": None, "hourly_rate": "8.5000", "funded_pension_percent": None,
        "fte": None, "employment_start_date": None, "employment_end_date": None,
    }]))
    assert employees["Employees"].max_row == 2
    assert employees["Employees"][2][3].value == "EE"

    returns = workbook(vat_returns_workbook(organisation, [{
        "id": "vat-1", "country_code": "EE", "period_start": "2026-01-01",
        "period_end": "2026-01-31", "basis": "accrual", "status": "Draft",
        "boxes_json": {"kmd_4_output_vat": "24.00"}, "reviewed_by": None,
        "reviewed_at": None, "submission_reference": None, "created_at": "2026-02-01T00:00:00Z",
    }]))
    assert [cell.value for cell in returns["VAT Returns"][1]] == [
        "id", "organisation_id", "country_code", "period_start", "period_end", "basis",
        "status", "boxes_json", "reviewed_by", "reviewed_at", "submission_reference", "created_at",
    ]
    assert '"kmd_4_output_vat": "24.00"' in returns["VAT Returns"][2][7].value


def test_payroll_service_workbook_draft_and_posted(db, ee_org):
    service = PayrollService(db)
    service.save_employee(ee_org["id"], actor=ACTOR, name="Workbook Employee", gross_salary="2000")
    draft = service.create_run(ee_org["id"], period="2026-07", actor=ACTOR)
    draft_book = workbook(service.workbook_xlsx(ee_org["id"], draft["id"]))
    assert draft_book.sheetnames == ["Summary", "Items"]

    posted = service.approve(ee_org["id"], draft["id"], actor=ACTOR)
    posted_book = workbook(service.workbook_xlsx(ee_org["id"], posted["id"]))
    assert posted_book.sheetnames == ["Summary", "Items", "Accounting"]
    accounting = posted_book["Accounting"]
    assert accounting.max_row > 3
    assert Decimal(str(accounting.cell(accounting.max_row, 4).value)) == Decimal(
        str(accounting.cell(accounting.max_row, 5).value)
    )


def test_payroll_service_workbook_is_tenant_scoped(db, ee_org, uk_org):
    service = PayrollService(db)
    service.save_employee(ee_org["id"], actor=ACTOR, name="Scoped Employee", gross_salary="2000")
    run = service.create_run(ee_org["id"], period="2026-07", actor=ACTOR)
    with pytest.raises(KeyError, match="Pay run not found"):
        service.workbook_xlsx(uk_org["id"], run["id"])


def test_excel_api_routes_authorisation_downloads_and_dates(db, ee_org, monkeypatch):
    service = PayrollService(db)
    service.save_employee(ee_org["id"], actor=ACTOR, name="API Workbook Employee", gross_salary="2000")
    run = service.create_run(ee_org["id"], period="2026-07", actor=ACTOR)
    monkeypatch.setenv("FASTACCOUNTS_DB", db.path)
    monkeypatch.setenv("DB_URL", "")
    monkeypatch.setenv("FASTACCOUNTS_ALLOW_TEST_AUTH", "true")
    reset_database_cache()
    from api_app import api

    root = f"/organisations/{ee_org['id']}"
    owner = {"X-Test-User": ACTOR}
    intruder = {"X-Test-User": "intruder@example.test"}
    with TestClient(api) as client:
        payroll_response = client.get(root + f"/pay-runs/{run['id']}/workbook.xlsx", headers=owner)
        assert payroll_response.status_code == 200
        assert "attachment" in payroll_response.headers["content-disposition"]
        assert workbook(payroll_response.content).sheetnames == ["Summary", "Items"]
        assert client.get(root + f"/pay-runs/{run['id']}/workbook.xlsx", headers=intruder).status_code == 403

        for report in ("general_ledger", "invoices", "employees", "vat_returns"):
            response = client.get(root + f"/reports/{report}.xlsx", headers=owner)
            assert response.status_code == 200, response.text
            assert response.headers["content-disposition"].endswith(f'"{report}.xlsx"')
            assert response.content.startswith(b"PK")
        assert client.get(root + "/reports/missing.xlsx", headers=owner).status_code == 404
        assert client.get(root + "/reports/invoices.xlsx", headers=intruder).status_code == 403
        assert client.get(
            root + "/reports/general_ledger.xlsx?date_from=2026-7-01", headers=owner
        ).status_code == 422
        assert client.get(
            root + "/reports/general_ledger.xlsx?date_from=2026-08-01&date_to=2026-07-01", headers=owner
        ).status_code == 422
    reset_database_cache()
