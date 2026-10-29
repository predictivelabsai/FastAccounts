"""Reviewed CSV/XLSX imports for tenant-scoped payroll dimensions and projects."""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
import json
from pathlib import Path
import zipfile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from core_utils import audit, canonical_json, new_id, utc_now
from database import Database, get_database

MAX_BYTES = 2_000_000
MAX_ROWS = 10_000
KINDS = {'department', 'position', 'cost_centre', 'project'}
FIELDS = {'code', 'name', 'group_code', 'group_name'}
REQUIRED = {
    'department': {'code', 'name'},
    'cost_centre': {'code', 'name'},
    'project': {'code', 'name'},
    'position': {'code', 'name', 'group_code', 'group_name'},
}
ALIASES = {'group': 'group_name', 'position_group': 'group_name', 'position_group_code': 'group_code'}


class PayrollDimensionImportService:
    def __init__(self, db: Database | None = None):
        self.db = db or get_database()

    @staticmethod
    def _read(file_name, content):
        suffix = Path(file_name).suffix.lower()
        if suffix not in {'.csv', '.xlsx'}:
            raise ValueError('Choose a CSV or XLSX file')
        if isinstance(content, str):
            raw = content.encode('utf-8')
        elif isinstance(content, bytes):
            raw = content
        else:
            raise ValueError('File content must be text or bytes')
        if not raw or len(raw) > MAX_BYTES:
            raise ValueError('Supply a non-empty file within the 2000000 byte limit')
        if suffix == '.csv':
            try:
                reader = csv.reader(StringIO(raw.decode('utf-8-sig'), newline=''), strict=True)
                records = list(reader)
            except (UnicodeDecodeError, csv.Error) as exc:
                raise ValueError('Invalid CSV file') from exc
            return records
        try:
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False)
            sheet = workbook.active
            records = []
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    if cell.data_type == 'f':
                        values.append('__FORMULA_CELL__')
                    else:
                        values.append('' if cell.value is None else str(cell.value).strip())
                records.append(values)
            workbook.close()
            return records
        except (InvalidFileException, OSError, zipfile.BadZipFile) as exc:
            raise ValueError('Invalid XLSX file') from exc

    @staticmethod
    def _columns(records, kind):
        if not records:
            raise ValueError('Import file has no header row')
        headers = [ALIASES.get(str(value or '').strip().lower().replace(' ', '_'),
                               str(value or '').strip().lower().replace(' ', '_')) for value in records[0]]
        if len(set(headers)) != len(headers) or not REQUIRED[kind] <= set(headers) or not set(headers) <= FIELDS:
            raise ValueError('File columns do not match the supported columns for ' + kind)
        return headers

    @staticmethod
    def _existing(tx, organisation_id, kind, fields):
        if kind == 'project':
            return tx.one("SELECT id FROM objects WHERE organisation_id=? AND object_type='project' AND code=?",
                          (organisation_id, fields['code']))
        return tx.one('SELECT id FROM payroll_dimensions WHERE organisation_id=? AND kind=? AND code=?',
                      (organisation_id, kind, fields['code']))

    @staticmethod
    def _validate(tx, organisation_id, kind, fields):
        errors = []
        if any(value == '__FORMULA_CELL__' for value in fields.values()):
            errors.append('Formula cells are not allowed in import files')
        for key, limit in (('code', 40), ('name', 180), ('group_code', 40), ('group_name', 180)):
            value = fields.get(key, '').strip()
            fields[key] = value
            if key in REQUIRED[kind] and not value:
                errors.append('Required field: ' + key.replace('_', ' '))
            if len(value) > limit:
                errors.append(key.replace('_', ' ').title() + ' is too long')
        if errors:
            return errors
        if kind == 'position':
            group = tx.one("SELECT id,name FROM payroll_dimensions WHERE organisation_id=? AND kind='position_group' AND code=?",
                           (organisation_id, fields['group_code']))
            if group and group['name'].strip().casefold() != fields['group_name'].casefold():
                errors.append('Position group code already has a different name')
        return errors

    def stage(self, organisation_id, kind, content, *, file_name, actor):
        if kind not in KINDS:
            raise ValueError('Unsupported payroll dimension import kind')
        records = self._read(file_name, content)
        headers = self._columns(records, kind)
        batch_id, now = new_id(), utc_now()
        seen = set()
        with self.db.transaction() as tx:
            if tx.execute('UPDATE organisations SET name=name WHERE id=?', (organisation_id,)).rowcount != 1:
                raise KeyError('Organisation not found')
            tx.execute('INSERT INTO payroll_dimension_import_batches(id,organisation_id,kind,status,created_by,created_at,updated_at) VALUES (?,?,? ,\'staged\',?,?,?)',
                       (batch_id, organisation_id, kind, actor, now, now))
            for row_number, values in enumerate(records[1:], 2):
                if not any(str(value or '').strip() for value in values):
                    continue
                if row_number > MAX_ROWS + 1:
                    raise ValueError('File exceeds the 10000 row limit')
                fields = {header: str(value or '').strip() for header, value in zip(headers, values)}
                error = 'File row has the wrong number of columns' if len(values) != len(headers) else ''
                errors = self._validate(tx, organisation_id, kind, fields)
                if error or errors:
                    error = '; '.join(filter(None, [error, *errors]))
                natural_key = (kind, fields.get('code', '').casefold())
                warning = ''
                if not error:
                    if natural_key in seen or self._existing(tx, organisation_id, kind, fields):
                        warning = 'Existing record will be skipped'
                    else:
                        seen.add(natural_key)
                tx.execute('INSERT INTO payroll_dimension_import_rows(id,organisation_id,batch_id,row_number,raw_json,parsed_json,error,warning,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?)',
                           (new_id(), organisation_id, batch_id, row_number, canonical_json(values),
                            canonical_json(fields), error, warning, now, now))
            self._counts(tx, organisation_id, batch_id)
            audit(tx, organisation_id, actor, 'payroll.dimension_import.staged', 'payroll_dimension_import', batch_id)
        return self.get(organisation_id, batch_id)

    @staticmethod
    def _counts(tx, organisation_id, batch_id):
        tx.execute('UPDATE payroll_dimension_import_batches SET total_rows=(SELECT COUNT(*) FROM payroll_dimension_import_rows WHERE organisation_id=? AND batch_id=?),'
                   'error_count=(SELECT COUNT(*) FROM payroll_dimension_import_rows WHERE organisation_id=? AND batch_id=? AND error<>\'\'),'
                   'warning_count=(SELECT COUNT(*) FROM payroll_dimension_import_rows WHERE organisation_id=? AND batch_id=? AND warning<>\'\'),updated_at=? '
                   'WHERE organisation_id=? AND id=?',
                   (organisation_id, batch_id, organisation_id, batch_id, organisation_id, batch_id,
                    utc_now(), organisation_id, batch_id))

    def get(self, organisation_id, batch_id, *, limit=100, offset=0):
        if not 1 <= limit <= MAX_ROWS or offset < 0:
            raise ValueError('Invalid import pagination')
        with self.db.transaction() as tx:
            batch = tx.one('SELECT * FROM payroll_dimension_import_batches WHERE organisation_id=? AND id=?',
                           (organisation_id, batch_id))
            if not batch:
                raise KeyError('Payroll dimension import batch not found')
            rows = tx.rows('SELECT * FROM payroll_dimension_import_rows WHERE organisation_id=? AND batch_id=? ORDER BY row_number LIMIT ? OFFSET ?',
                           (organisation_id, batch_id, limit, offset))
            batch['rows'] = [{**row, 'raw': json.loads(row['raw_json']), 'parsed': json.loads(row['parsed_json'])} for row in rows]
            batch['limit'], batch['offset'] = limit, offset
            return batch

    def commit(self, organisation_id, batch_id, *, actor):
        with self.db.transaction() as tx:
            if tx.execute('UPDATE organisations SET name=name WHERE id=?', (organisation_id,)).rowcount != 1:
                raise KeyError('Organisation not found')
            batch = tx.one('SELECT * FROM payroll_dimension_import_batches WHERE organisation_id=? AND id=?',
                           (organisation_id, batch_id))
            if not batch:
                raise KeyError('Payroll dimension import batch not found')
            if batch['status'] != 'staged':
                raise ValueError('Only a staged batch can be committed or discarded')
            tx.execute("UPDATE payroll_dimension_import_batches SET status='committed',updated_at=? WHERE organisation_id=? AND id=?",
                       (utc_now(), organisation_id, batch_id))
            rows = tx.rows('SELECT * FROM payroll_dimension_import_rows WHERE organisation_id=? AND batch_id=? ORDER BY row_number',
                           (organisation_id, batch_id))
            for row in rows:
                if row['error']:
                    continue
                fields = json.loads(row['parsed_json'])
                if self._existing(tx, organisation_id, batch['kind'], fields):
                    self._outcome(tx, organisation_id, row['id'], warning='Existing record skipped')
                    continue
                errors = self._validate(tx, organisation_id, batch['kind'], fields)
                if errors:
                    self._outcome(tx, organisation_id, row['id'], error='; '.join(errors))
                    continue
                if batch['kind'] == 'project':
                    entity_id, now = new_id(), utc_now()
                    tx.execute("INSERT INTO objects(id,organisation_id,object_type,active,code,name,notes,created_at,updated_at) VALUES (?,?,'project',1,?,?,'',?,?)",
                               (entity_id, organisation_id, fields['code'], fields['name'], now, now))
                else:
                    parent_id = None
                    if batch['kind'] == 'position':
                        parent = tx.one("SELECT id FROM payroll_dimensions WHERE organisation_id=? AND kind='position_group' AND code=?",
                                        (organisation_id, fields['group_code']))
                        if parent:
                            parent_id = parent['id']
                        else:
                            parent_id, now = new_id(), utc_now()
                            tx.execute("INSERT INTO payroll_dimensions(id,organisation_id,kind,code,name,parent_id,active,created_at,updated_at) VALUES (?,?,'position_group',?,?,NULL,1,?,?)",
                                       (parent_id, organisation_id, fields['group_code'], fields['group_name'], now, now))
                    entity_id, now = new_id(), utc_now()
                    tx.execute('INSERT INTO payroll_dimensions(id,organisation_id,kind,code,name,parent_id,active,created_at,updated_at) VALUES (?,?,?,?,?,?,1,?,?)',
                               (entity_id, organisation_id, batch['kind'], fields['code'], fields['name'], parent_id, now, now))
                self._outcome(tx, organisation_id, row['id'], entity_id=entity_id)
            self._counts(tx, organisation_id, batch_id)
            audit(tx, organisation_id, actor, 'payroll.dimension_import.committed', 'payroll_dimension_import', batch_id)
        return self.get(organisation_id, batch_id)

    @staticmethod
    def _outcome(tx, organisation_id, row_id, *, entity_id=None, error='', warning=''):
        tx.execute('UPDATE payroll_dimension_import_rows SET entity_id=?,error=?,warning=?,updated_at=? WHERE organisation_id=? AND id=?',
                   (entity_id, error, warning, utc_now(), organisation_id, row_id))

    def discard(self, organisation_id, batch_id, *, actor):
        with self.db.transaction() as tx:
            batch = tx.one('SELECT status FROM payroll_dimension_import_batches WHERE organisation_id=? AND id=?',
                           (organisation_id, batch_id))
            if not batch:
                raise KeyError('Payroll dimension import batch not found')
            if batch['status'] != 'staged':
                raise ValueError('Only a staged batch can be committed or discarded')
            tx.execute("UPDATE payroll_dimension_import_batches SET status='discarded',updated_at=? WHERE organisation_id=? AND id=?",
                       (utc_now(), organisation_id, batch_id))
            audit(tx, organisation_id, actor, 'payroll.dimension_import.discarded', 'payroll_dimension_import', batch_id)
        return self.get(organisation_id, batch_id)
